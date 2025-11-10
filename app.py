from flask import Flask, render_template, request, send_file
import openai
from pptx import Presentation
import os
import json

app = Flask(__name__)

# Cấu hình API key của OpenAI (lưu ý: nên dùng biến môi trường)
# openai.api_key = os.getenv("OPENAI_API_KEY", "YOUR_API_KEY")

def generate_slides_from_ai(topic, data):
    """
    Gửi yêu cầu đến API của OpenAI để tạo nội dung slide.
    Hàm này sẽ bị comment lại vì không có API key.
    """
    # prompt = f"""
    # Hãy tạo nội dung cho một bài trình chiếu PowerPoint dựa trên chủ đề sau: '{topic}' và dữ liệu được cung cấp.
    # Dữ liệu:
    # {data}

    # Vui lòng trả về kết quả dưới dạng một đối tượng JSON hợp lệ.
    # Cấu trúc JSON phải là một danh sách các slide, trong đó mỗi slide là một đối tượng có 'title' và 'content'.
    # 'content' phải là một danh sách các chuỗi (string), mỗi chuỗi là một gạch đầu dòng.
    # Ví dụ:
    # {{
    #   "slides": [
    #     {{
    #       "title": "Slide Tiêu đề",
    #       "content": ["Nội dung gạch đầu dòng 1", "Nội dung gạch đầu dòng 2"]
    #     }},
    #     {{
    #       "title": "Slide thứ hai",
    #       "content": ["Điểm chính 1", "Điểm chính 2", "Điểm chính 3"]
    #     }}
    #   ]
    # }}
    # """
    # response = openai.Completion.create(
    #     engine="text-davinci-003", # Hoặc một model mới hơn
    #     prompt=prompt,
    #     max_tokens=1500,
    #     n=1,
    #     stop=None,
    #     temperature=0.7,
    # )
    # return json.loads(response.choices[0].text)
    pass

def generate_slides_from_ai_mock(topic, data):
    """
    Hàm giả lập để trả về cấu trúc JSON mẫu mà không cần gọi API thật.
    """
    mock_response = {
        "slides": [
            {
                "title": f"Giới thiệu về {topic}",
                "content": [
                    "Đây là slide giới thiệu tổng quan.",
                    "Dữ liệu đầu vào sẽ được phân tích và trình bày.",
                    f"Dữ liệu cung cấp: {data[:100]}..."
                ]
            },
            {
                "title": "Phân tích Các điểm chính",
                "content": [
                    "Điểm chính thứ nhất từ dữ liệu.",
                    "Điểm chính thứ hai, làm rõ hơn.",
                    "Hệ quả và những điều cần lưu ý."
                ]
            },
            {
                "title": "Kết luận",
                "content": [
                    "Tóm tắt các phát hiện chính.",
                    "Đề xuất các bước tiếp theo.",
                    "Cảm ơn đã lắng nghe."
                ]
            }
        ]
    }
    return mock_response


def create_presentation(slides_data):
    """
    Tạo một file PowerPoint từ dữ liệu slide có cấu trúc.
    """
    prs = Presentation()
    title_slide_layout = prs.slide_layouts[0] # Layout cho slide tiêu đề
    content_slide_layout = prs.slide_layouts[1] # Layout cho slide nội dung

    for i, slide_info in enumerate(slides_data["slides"]):
        if i == 0: # Slide đầu tiên thường là slide tiêu đề chính
            slide = prs.slides.add_slide(title_slide_layout)
            title = slide.shapes.title
            subtitle = slide.placeholders[1]
            title.text = slide_info["title"]
            if slide_info["content"]:
                subtitle.text = "\n".join(slide_info["content"])
        else:
            slide = prs.slides.add_slide(content_slide_layout)
            title = slide.shapes.title
            body = slide.shapes.placeholders[1]
            title.text = slide_info["title"]

            # Thêm nội dung gạch đầu dòng
            tf = body.text_frame
            tf.clear() # Xóa placeholder text mặc định
            for line in slide_info["content"]:
                p = tf.add_paragraph()
                p.text = line
                p.level = 1

    return prs


@app.route('/')
def index():
    return render_template('index.html')

@app.route('/generate', methods=['POST'])
def generate():
    topic = request.form['topic']
    data = request.form.get('data', '') # Sử dụng .get để tránh KeyError
    file = request.files.get('file')

    if file and file.filename != '':
        # Ưu tiên đọc dữ liệu từ file nếu được cung cấp
        if file.filename.endswith('.txt'):
            data = file.read().decode('utf-8')
        else:
            return "Lỗi: Chỉ hỗ trợ file .txt", 400
    elif not data:
        return "Lỗi: Vui lòng cung cấp dữ liệu qua văn bản hoặc tải lên một file.", 400

    # Gọi hàm AI để lấy nội dung (sử dụng hàm giả lập)
    slides_content = generate_slides_from_ai_mock(topic, data)

    # Tạo file PowerPoint
    presentation = create_presentation(slides_content)

    # Lưu file vào một buffer trong bộ nhớ
    import io
    file_stream = io.BytesIO()
    presentation.save(file_stream)
    file_stream.seek(0)

    # Gửi file về cho người dùng
    return send_file(
        file_stream,
        as_attachment=True,
        download_name=f"{topic.replace(' ', '_')}_presentation.pptx"
    )

if __name__ == '__main__':
    app.run(debug=True)
